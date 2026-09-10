"""Generate editorial HTML, DOCX and custom figures for both analyses."""
from pathlib import Path
from urllib.parse import quote
import argparse, html, re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from rebuild_analysis_sources import rebuild

ROOT=Path(__file__).resolve().parents[1]
FONT=Path("C:/Windows/Fonts/msyh.ttc"); FONT_B=Path("C:/Windows/Fonts/msyhbd.ttc")
SPECS={
 "delta":dict(folder="02 delta-operator-research",title="三角洲行动烽火地带玩法系统分析",subtitle="从一套配装到一次撤离 干员小队如何进入搜打撤决策",accent="#3f6f62",soft="#e8f0ed",html="三角洲行动烽火地带干员系统研究.html",pdf="三角洲行动烽火地带干员系统研究.pdf",docx="三角洲行动玩法分析.docx"),
 "valorant":dict(folder="03 valorant-agent-research",title="VALORANT战术系统分析",subtitle="从停步开枪到五人进点 英雄技能怎样改变信息 空间与节奏",accent="#9b3747",soft="#f4e9eb",html="VALORANT英雄技能系统研究.html",pdf="VALORANT英雄技能系统研究.pdf",docx="VALORANT玩法分析.docx")}


class Fig:
 def __init__(self,path,spec,title):
  self.path=path; self.a=spec["accent"]; self.s=spec["soft"]; self.w=1600; self.h=900
  self.im=Image.new("RGB",(self.w,self.h),"#fbfaf7"); self.d=ImageDraw.Draw(self.im)
  self.svg=[f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {self.w} {self.h}" role="img" aria-label="{html.escape(title)}"><rect width="100%" height="100%" fill="#fbfaf7"/>']
 def text(self,x,y,s,size=28,color="#202628",bold=False,anchor="la"):
  self.d.text((x,y),s,font=ImageFont.truetype(str(FONT_B if bold else FONT),size),fill=color,anchor=anchor)
  self.svg.append(f'<text x="{x}" y="{y}" font-family="Microsoft YaHei,sans-serif" font-size="{size}" font-weight="{700 if bold else 400}" fill="{color}" dominant-baseline="middle" text-anchor="{"middle" if anchor=="mm" else "start"}">{html.escape(s)}</text>')
 def rect(self,x1,y1,x2,y2,fill="#fff",outline="#cbd2d0",r=15,w=2):
  self.d.rounded_rectangle((x1,y1,x2,y2),radius=r,fill=fill,outline=outline,width=w)
  self.svg.append(f'<rect x="{x1}" y="{y1}" width="{x2-x1}" height="{y2-y1}" rx="{r}" fill="{fill}" stroke="{outline}" stroke-width="{w}"/>')
 def line(self,pts,color=None,w=5,arrow=False,dash=False):
  color=color or self.a; self.d.line(pts,fill=color,width=w)
  self.svg.append(f'<polyline points="{" ".join(f"{x},{y}" for x,y in pts)}" fill="none" stroke="{color}" stroke-width="{w}"'+(' stroke-dasharray="12 10"' if dash else '')+'/>' )
  if arrow:
   import math
   x,y=pts[-1]; px,py=pts[-2]; a=math.atan2(y-py,x-px)
   tri=[(x,y),(x-22*math.cos(a-.5),y-22*math.sin(a-.5)),(x-22*math.cos(a+.5),y-22*math.sin(a+.5))]
   self.d.polygon(tri,fill=color); self.svg.append('<polygon points="'+' '.join(f'{x},{y}' for x,y in tri)+f'" fill="{color}"/>')
 def save(self):
  self.path.parent.mkdir(parents=True,exist_ok=True); self.im.save(self.path,optimize=True)
  self.path.with_suffix(".svg").write_text("".join(self.svg)+"</svg>",encoding="utf-8")


def row(f,labels,y,subs):
 gap=30; margin=60; ww=(f.w-2*margin-gap*(len(labels)-1))//len(labels); centers=[]
 for i,(lab,sub) in enumerate(zip(labels,subs)):
  x=margin+i*(ww+gap); f.rect(x,y,x+ww,y+130); f.text(x+ww//2,y+45,lab,29,f.a,True,"mm"); f.text(x+ww//2,y+91,sub,19,"#5d686a",False,"mm"); centers.append((x+ww//2,y+65,ww))
 return centers


def figures(kind,s):
 out=ROOT/s["folder"]/"assets"
 if kind=="delta":
  f=Fig(out/"delta-system.png",s,"烽火地带局内外系统全景"); f.text(70,70,"局外准备决定带什么进去 局内撤离决定什么能够留下",40,"#172124",True)
  a=row(f,["任务与设施","仓库与交易","配装预算"],160,["材料需求与阶段目标","库存 购买 出售","枪 弹 甲 药 容量"]); b=row(f,["搜索与任务","交战与干员","路线与撤离"],445,["提高收益 改写目标","改变信息 位置 状态","兑现收益或承担损失"])
  for x,y,w in a:f.line([(x,y+65),(x,400)],w=4,arrow=True)
  for x,y,w in b:f.line([(x,y+65),(x,755)],w=4,arrow=True)
  f.line([(140,755),(1460,755)],"#82908d",3); f.text(800,815,"结算回到仓库 再形成下一局需求",27,"#394648",True,"mm"); f.save()
  f=Fig(out/"delta-loop.png",s,"从配装到结算的风险收益循环"); f.text(70,70,"每一步都在重估收益 成本与撤离概率",40,"#172124",True)
  c=row(f,["战前配装","进入战区","搜索取舍","接战或绕行","撤离判定","资产结算"],310,["确定成本 C","信息不完整","提高 V 与 G","改变 p 与消耗","出口与时机","成功保留 失败损失"])
  for q,n in zip(c,c[1:]): f.line([(q[0]+q[2]//2,q[1]),(n[0]-n[2]//2,n[1])],w=5,arrow=True)
  f.text(800,650,"价值越高 不一定越想继续打 技能也可能被用于避战",28,s["accent"],True,"mm"); f.save()
  f=Fig(out/"delta-squad.png",s,"三人小队的任务缺口与工具补位"); f.text(70,70,"先列局面要完成的动作 再决定谁来补工具",40,"#172124",True)
  data=[("确认枪线","侦察 视野 探头","露娜给局部线索 队友仍要处理未知角"),("穿过空地","烟幕 机动 分段通过","红狼帮助突破 但不消除第二条枪线"),("战后恢复","治疗 药品 安全位置","蜂医补生命值 不补护甲 弹药与站位"),("照看后路","装置 留人 声音信息","牧羊人增加处理步骤 对手仍可识别 绕行")]
  for i,(a,b,c) in enumerate(data):
   y=155+i*165; f.rect(90,y,485,y+110,s["soft"],s["accent"]); f.text(287,y+38,a,29,s["accent"],True,"mm"); f.text(287,y+79,b,19,"#526062",False,"mm"); f.line([(485,y+55),(620,y+55)],w=5,arrow=True); f.rect(620,y,1510,y+110); f.text(660,y+55,c,24)
  f.save()
  f=Fig(out/"delta-route.png",s,"满包撤离的两条假设路线"); f.text(70,70,"假设地形 不对应真实地图点位",34,"#172124",True)
  f.rect(610,210,930,650,"#ccd5cf","#aab6b0"); f.text(770,430,"建筑遮挡",30,"#43504d",True,"mm"); f.rect(1160,170,1470,280,"#f0dddd","#bb7777"); f.text(1315,225,"疑似敌方枪线",25,"#8c4545",True,"mm"); f.rect(1210,650,1500,770,s["soft"],s["accent"]); f.text(1355,710,"撤离方向",30,s["accent"],True,"mm")
  f.rect(90,170,350,290); f.text(220,215,"满包三人小队",28,s["accent"],True,"mm"); f.text(220,255,"一人受伤",22,"#5d686a",False,"mm"); f.line([(350,220),(1080,220),(1260,610)],w=7,arrow=True); f.text(750,165,"A 短但暴露",26,s["accent"],True,"mm"); f.line([(220,290),(220,720),(1160,720)],w=7,arrow=True,dash=True); f.text(650,775,"B 遮挡多 但出口和耗时未知",26,s["accent"],True,"mm"); f.line([(1180,275),(930,420)],"#a74c4c",5,True,True); f.save()
 else:
  f=Fig(out/"valorant-system.png",s,"标准爆破的系统全景"); f.text(70,70,"英雄技能建立在射击 地图 经济和回合目标之上",40,"#172124",True)
  for k,(labs,tag) in enumerate([(["移动与停步","武器与护甲","声音与沟通"],"操作底盘"),(["地图枪线","区域控制","爆能器与时间"],"空间目标"),(["英雄选择","技能资源","队伍分工"],"角色工具")]): row(f,labs,145+k*225,[tag]*3)
  f.text(800,835,"系统共同决定能看见什么 能走到哪里 何时必须行动",25,s["accent"],True,"mm"); f.save()
  f=Fig(out/"valorant-cycle.png",s,"信息到执行的回合循环"); f.text(70,70,"技能不是终点 反馈必须重新进入计划",40,"#172124",True)
  pts=[(300,230),(1100,230),(1100,620),(300,620)]; labs=[("信息","敌位 枪线 资源"),("计划","处理顺序与分工"),("执行","技能 移动 射击"),("反馈","击杀 退让 反制")]
  for (x,y),(a,b) in zip(pts,labs): f.rect(x-190,y-75,x+190,y+75,"#fff",s["accent"]); f.text(x,y-20,a,34,s["accent"],True,"mm"); f.text(x,y+30,b,22,"#5d686a",False,"mm")
  for a,b in [((490,230),(900,230)),((1100,305),(1100,545)),((910,620),(490,620)),((300,545),(300,305))]:f.line([a,b],w=7,arrow=True)
  f.text(800,430,"Intel → Plan → Execute → Repeat",30,"#283337",True,"mm"); f.save()
  f=Fig(out/"valorant-timeline.png",s,"三段条件的进点重叠窗口"); f.text(70,70,"假设秒数 只解释时序 不代表现行技能参数",34,"#172124",True); left,right=300,1500
  for t in range(11): x=left+(right-left)*t/10; f.line([(x,170),(x,700)],"#d7dcda",2); f.text(x,145,str(t),18,"#6a7475",False,"mm")
  for i,(name,a,b,c) in enumerate([("侦察信息",2,8,"#577d89"),("远线遮挡",3,7,"#8a6c91"),("队友跟进",4,6,s["accent"])]):
   y=230+i*145; f.text(80,y+45,name,26,"#283337",True); x1=left+(right-left)*a/10; x2=left+(right-left)*b/10; f.rect(x1,y,x2,y+80,c,c,12); f.text((x1+x2)/2,y+40,f"{a} 至 {b} 秒",22,"#fff",True,"mm")
  f.rect(780,680,1020,760,"#d3a34d","#d3a34d",10); f.text(900,720,"共同窗口 2 秒",23,"#fff",True,"mm"); f.save()
  f=Fig(out/"valorant-postplant.png",s,"安装前后目标与时间压力反转"); f.text(70,70,"地图没变 必须主动解决问题的一方变了",40,"#172124",True)
  for x,title,fill in [(80,"安装前","#fff"),(860,"安装后",s["soft"])]: f.rect(x,180,x+660,680,fill,s["accent"]); f.text(x+330,230,title,35,s["accent"],True,"mm")
  for x,a,b,c,d in [(145,"进攻方","进入安装区 清枪线 付出时间","防守方","守住入口 延缓并消耗对手"),(925,"防守方","回到包点 拆除 承担时间压力","进攻方","守包 拖延 迫使对方主动进入")]: f.text(x,325,a,28,"#273234",True); f.text(x,375,b,23,"#596466"); f.text(x,490,c,28,"#273234",True); f.text(x,540,d,23,"#596466")
  f.line([(740,430),(850,430)],w=8,arrow=True); f.save()


def parse(path):
 lines=path.read_text(encoding="utf-8").splitlines(); out=[]; i=0
 while i<len(lines):
  s=lines[i].rstrip()
  if not s:i+=1;continue
  if s.startswith("|"):
   rows=[]
   while i<len(lines) and lines[i].startswith("|"): rows.append([c.strip() for c in lines[i].strip("|").split("|")]);i+=1
   if len(rows)>1 and all(re.fullmatch(r"[-: ]+",c or "-") for c in rows[1]):rows.pop(1)
   out.append(("table",rows));continue
  if s.startswith("!["):
   out.append(("image",re.match(r"!\[(.*?)\]\((.*?)\)",s).groups()));i+=1;continue
  m=re.match(r"^(#{1,3})\s+(.*)",s)
  if m:out.append(("heading",(len(m[1]),m[2])));i+=1;continue
  if re.match(r"^\d+\. ",s):out.append(("list",s));i+=1;continue
  if s.startswith("- "):out.append(("bullet",s[2:]));i+=1;continue
  p=[s];i+=1
  while i<len(lines) and lines[i].strip() and not re.match(r"^(#{1,3})\s+|^\d+\. |^- |^\||^!\[",lines[i]):p.append(lines[i].strip());i+=1
  out.append(("para"," ".join(p)))
 return out


def inline_doc(p,text):
 pos=0
 for m in re.finditer(r"https?://[^\s）]+",text):
  p.add_run(text[pos:m.start()]); rid=p.part.relate_to(m.group(),"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True); link=OxmlElement("w:hyperlink");link.set(qn("r:id"),rid);run=OxmlElement("w:r");t=OxmlElement("w:t");t.text=m.group();run.append(t);link.append(run);p._p.append(link);pos=m.end()
 p.add_run(text[pos:])


def cell(c,text,header,accent):
 c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;p=c.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(0);r=p.add_run(text);r.bold=header;r.font.size=Pt(8.5);r.font.color.rgb=RGBColor.from_string("FFFFFF" if header else "222222")
 shd=OxmlElement("w:shd");shd.set(qn("w:fill"),accent.replace("#","") if header else "F5F6F5");c._tc.get_or_add_tcPr().append(shd)
 mar=OxmlElement("w:tcMar")
 for side in ("top","left","bottom","right"):e=OxmlElement("w:"+side);e.set(qn("w:w"),"100");e.set(qn("w:type"),"dxa");mar.append(e)
 c._tc.get_or_add_tcPr().append(mar)


def docx(s,blocks):
 d=Document();sec=d.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297);sec.top_margin=sec.bottom_margin=Mm(19);sec.left_margin=sec.right_margin=Mm(23)
 for name,font,size in [("Normal","宋体",10.5),("Title","黑体",24),("Heading 1","黑体",16),("Heading 2","黑体",12.5)]:st=d.styles[name];st.font.name="Times New Roman";st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),font)
 for b in d.styles.element.findall(".//"+qn("w:pBdr")):b.getparent().remove(b)
 n=d.styles["Normal"].paragraph_format;n.line_spacing=1.45;n.space_after=Pt(6);n.first_line_indent=Pt(21);n.widow_control=True
 for nm in ("Heading 1","Heading 2"):p=d.styles[nm].paragraph_format;p.space_before=Pt(15);p.space_after=Pt(7);p.keep_with_next=True;p.first_line_indent=Pt(0)
 foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");foot._p.append(fld)
 d.core_properties.title=s["title"];d.core_properties.subject=s["subtitle"];d.core_properties.author="";d.core_properties.last_modified_by=""
 p=d.add_paragraph(style="Title");p.paragraph_format.space_before=Pt(58);p.add_run(s["title"]);p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.add_run(s["subtitle"]).font.size=Pt(13);p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(18);p.add_run("玩法系统研究  2026 年 9 月").font.size=Pt(10);d.add_page_break();d.add_heading("目录",1)
 for typ,data in blocks:
  if typ=="heading" and data[0]==2:p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(3);p.add_run(data[1])
 d.add_page_break();images=0
 for typ,data in blocks:
  if typ=="heading":
   if data[0]>1:d.add_heading(data[1],1 if data[0]==2 else 2)
  elif typ=="para":
   p=d.add_paragraph();inline_doc(p,data)
   if data.startswith(("图 ","资料核对","来源：")):p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.1
  elif typ in ("list","bullet"):
   p=d.add_paragraph(style="Normal" if typ=="list" else "List Bullet");p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Mm(7);inline_doc(p,data)
  elif typ=="image":
   p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pic=p.add_run().add_picture(str(ROOT/s["folder"]/data[1]),width=Mm(160));pic._inline.docPr.set("descr",data[0]);images+=1
  elif typ=="table":
   t=d.add_table(rows=len(data),cols=max(map(len,data)));t.autofit=False
   for i,r in enumerate(data):
    for j,v in enumerate(r):cell(t.cell(i,j),v,i==0,s["accent"])
   t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"));d.add_paragraph()
 out=ROOT/s["folder"]/s["docx"];d.save(out);return out,images


def inhtml(t):
 e=html.escape(t);return re.sub(r"(https?://[^\s）]+)",r'<a href="\1">\1</a>',e)


def webpage(s,blocks):
 body=[];toc=[]
 for typ,data in blocks:
  if typ=="heading":
   level,text=data
   if level==1:continue
   sid="sec-"+re.match(r"\d+",text).group() if level==2 and re.match(r"\d+",text) else ""
   if level==2:toc.append((sid,text))
   body.append(f'<h{level} id="{sid}">{html.escape(text)}</h{level}>')
  elif typ=="para":body.append("<p>"+inhtml(data)+"</p>")
  elif typ in ("list","bullet"):body.append(f'<p class="item"><span>•</span>{inhtml(data)}</p>')
  elif typ=="image":body.append(f'<figure><a href="{Path(data[1]).with_suffix(".svg").as_posix()}"><img src="{Path(data[1]).with_suffix(".svg").as_posix()}" alt="{html.escape(data[0])}"></a></figure>')
  elif typ=="table":body.append('<div class="table-wrap"><table><thead><tr>'+''.join('<th>'+inhtml(x)+'</th>' for x in data[0])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+inhtml(x)+'</td>' for x in r)+'</tr>' for r in data[1:])+'</tbody></table></div>')
 nav=''.join(f'<a href="#{i}"><span>{n:02d}</span>{html.escape(t[3:])}</a>' for n,(i,t) in enumerate(toc,1))
 css=f'''*{{box-sizing:border-box}}html{{scroll-behavior:smooth}}body{{margin:0;color:#202628;background:#f3f1ec;font:17px/1.9 "Songti SC","SimSun",serif}}a{{color:{s['accent']};text-underline-offset:3px}}.layout{{max-width:1260px;margin:auto;display:grid;grid-template-columns:270px minmax(0,850px);gap:54px;padding:54px 28px 100px}}aside{{position:sticky;top:22px;height:calc(100vh - 44px);overflow:auto;font:14px/1.5 "Microsoft YaHei",sans-serif}}aside h1{{font-size:20px;line-height:1.45;margin:0 0 12px}}aside p{{font-size:13px;color:#657071;line-height:1.7}}nav{{border-top:1px solid #cfd5d2;margin-top:22px;padding-top:12px}}nav a{{display:grid;grid-template-columns:28px 1fr;gap:8px;padding:7px 0;color:#344143;text-decoration:none}}nav span{{color:{s['accent']}}}main{{background:#fff;padding:64px 76px;box-shadow:0 14px 45px #35403c17}}.hero{{padding-bottom:42px;border-bottom:1px solid #d7dcda;margin-bottom:48px}}.eyebrow{{font:13px "Microsoft YaHei",sans-serif;color:{s['accent']};letter-spacing:.12em}}.hero h1{{font:700 37px/1.32 "Microsoft YaHei",sans-serif;margin:16px 0}}.hero p{{font-size:19px;color:#5a6566;text-indent:0}}.downloads a{{font:14px "Microsoft YaHei",sans-serif;margin-right:22px}}h2,h3{{font-family:"Microsoft YaHei",sans-serif;color:#172124;line-height:1.45}}h2{{font-size:26px;margin:58px 0 22px;padding-top:8px}}h3{{font-size:19px;margin:32px 0 12px}}p{{margin:0 0 17px;text-indent:2em}}.item{{text-indent:0;padding-left:25px;position:relative}}.item span{{position:absolute;left:0;color:{s['accent']}}}figure{{margin:34px -28px 12px}}figure img{{display:block;width:100%}}.table-wrap{{overflow:auto;margin:26px 0 34px}}table{{border-collapse:collapse;width:100%;font:14px/1.55 "Microsoft YaHei",sans-serif}}th,td{{border:1px solid #d8dddb;padding:12px 13px;text-align:left;vertical-align:top}}th{{background:{s['accent']};color:white}}tbody tr:nth-child(even){{background:#f5f6f5}}footer{{margin-top:60px;padding-top:22px;border-top:1px solid #d7dcda;font-size:13px;color:#687172}}@media(max-width:900px){{.layout{{display:block;padding:0}}aside{{position:static;height:auto;padding:25px}}aside nav{{display:none}}main{{padding:38px 22px;box-shadow:none}}figure{{margin:28px 0}}.hero h1{{font-size:30px}}}}@media print{{body{{background:#fff}}aside{{display:none}}.layout{{display:block;padding:0}}main{{box-shadow:none;padding:15mm}}.downloads{{display:none}}h2{{break-after:avoid}}figure,.table-wrap{{break-inside:avoid}}}}'''
 page=f'''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(s['title'])}</title><style>{css}</style></head><body><div class="layout"><aside><h1>{html.escape(s['title'])}</h1><p>{html.escape(s['subtitle'])}</p><nav>{nav}</nav></aside><main><header class="hero"><div class="eyebrow">GAMEPLAY SYSTEM STUDY · 2026</div><h1>{html.escape(s['title'])}</h1><p>{html.escape(s['subtitle'])}</p><div class="downloads"><a href="{quote(s['pdf'])}">下载 PDF</a><a href="{quote(s['docx'])}">下载 Word</a><a href="正文.md">查看正文</a><a href="信息来源汇总.md">信息来源</a></div></header>{''.join(body)}<footer>本文用于游戏策划作品集。规则事实、历史改动、本文推演与待验证方案分别标注。</footer></main></div></body></html>'''
 out=ROOT/s["folder"]/s["html"];out.write_text(page,encoding="utf-8");return out


def sources(kind,s):
 if kind=="delta":rows=[("D0","Garena Operations Mode 101","https://deltaforce.garena.com/en/news/system/RMDWVM","2026-02-20；模式循环、仓库、任务、配装、弹甲、声音与撤离。"),("D1","Garena 模式与红狼介绍","https://deltaforce.garena.com/en/?redirect=0","三人小队、四类职能与代表技能。"),("D2","Starfall 更新公告","https://deltaforce.garena.com/en/news/announcement/H3T8DS","2025-01；蜂医、露娜、威龙双模式历史案例。"),("D3","Break 更新公告","https://deltaforce.garena.com/en/news/announcement/V6J38U","2025-07-08；牧羊人陷阱提示与反制。"),("D4","干员平衡说明","https://deltaforce.garena.com/en/news/all/A7WYBP","2026-07-10；道具密度、单向信息与枪战可读性。")]
 else:rows=[("V1","Riot How we balance VALORANT","https://playvalorant.com/en-us/news/dev/how-we-balance-valorant/","2020-06-26；战术循环。"),("V2","Riot Beginner's Guide","https://playvalorant.com/en-us/news/announcements/beginners-guide/","2024-08-02；英雄、回合、经济、武器与地图。"),("V3","Sova 官方页","https://playvalorant.com/en-us/agents/sova/","侦察箭判定。"),("V4","Omen 官方页","https://playvalorant.com/en-us/agents/omen/","烟雾机制。"),("V5","Cypher 官方页","https://playvalorant.com/en-us/agents/cypher/","绊线机制。"),("V6","4.08 版本说明","https://playvalorant.com/en-us/news/game-updates/valorant-patch-notes-4-08/","Jett 历史调整。"),("V7","7.04 版本说明","https://playvalorant.com/en-us/news/game-updates/valorant-patch-notes-7-04/","Jett 历史调整。"),("V8","9.10 版本说明","https://playvalorant.com/en-us/news/game-updates/valorant-patch-notes-9-10/","Omen 烟雾几何约束。"),("V9","Valve CS2","https://www.counter-strike.net/cs2","Responsive Smokes 横向比较。")]
 text=f"# {s['title']} 信息来源\n\n资料核对：2026-09-10。历史参数不作当前版本攻略。\n\n| 编号 | 来源 | 日期与用途 |\n|---|---|---|\n"+''.join(f'| {k} | [{t}]({u}) | {d} |\n' for k,t,u,d in rows)
 text+="""\n## 证据边界\n\n- 官方资料只支撑相应机制事实；推演、评价与建议是本文判断。\n- 自绘图不是实机截图；假设秒数、概率、价值单位与路线不是实测。\n- 未采集用户实战、访谈、胜率或后台日志；待验证方案没有写成测试结果。\n- 历史地区公告只证明对应地区与日期，不自动证明国服当前参数。\n\n## 知识库与版式参考\n\n按需参考 game-system-design-kb 的 00 方法论、01 战斗系统和 03 经济系统。参考用户指定公开拆解的信息层级、目录导航和图文关系，没有复制其正文、截图或图表。\n"""
 (ROOT/s["folder"]/"信息来源汇总.md").write_text(text,encoding="utf-8")


def main():
 a=argparse.ArgumentParser();a.add_argument("--reset-source",action="store_true");args=a.parse_args()
 if args.reset_source:rebuild()
 for kind,s in SPECS.items():
  figures(kind,s);blocks=parse(ROOT/s["folder"]/"正文.md");d,n=docx(s,blocks);h=webpage(s,blocks);sources(kind,s);print(kind,len(blocks),n,d.name,h.name)

if __name__=="__main__":main()
